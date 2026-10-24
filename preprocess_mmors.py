"""
MMORS Water Quality Results (2012–2018) - Data Preprocessing
Run: python preprocess_mmors.py
Requirements: pandas, openpyxl, numpy
Place the original workbook beside this script or update INPUT_FILE.
"""
import pandas as pd
import numpy as np
import re
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

INPUT_FILE = "MMORS_water_quality_results_2012-2018_orig.xlsx"
OUTPUT_FILE = "MMORS_water_quality_preprocessed.xlsx"
SHEETS = ["Marilao", "Meycauayan", "Obando"]

METRIC_NAMES = [
    "Latitude_North_deg", "Longitude_East_deg", "Dissolved_Oxygen_mg_L", "pH",
    "Temperature_C", "Biochemical_Oxygen_Demand_mg_L",
    "Total_Suspended_Solids_mg_L", "Color_TCU",
    "Fecal_Coliform_MPN_100mL", "Total_Coliform_MPN_100mL",
    "Ammonia_mg_L", "Nitrates_as_Nitrogen_mg_L",
    "Phosphates_as_Phosphorous_mg_L", "Chlorides_mg_L"
]

def normalize_time(value):
    """Standardize common time formats without inventing unavailable times."""
    if pd.isna(value) or str(value).strip().lower() in ("nan", "nat", ""):
        return np.nan
    if hasattr(value, "strftime"):
        return value.strftime("%H:%M:%S")
    try:
        return pd.to_datetime(str(value).strip()).strftime("%H:%M:%S")
    except (ValueError, TypeError):
        return str(value).strip()

def extract_records(input_file):
    records = []
    extraction_summary = []
    no_sampling_total = 0

    for sheet in SHEETS:
        raw = pd.read_excel(input_file, sheet_name=sheet, header=None)
        current_month = None
        usable_count = 0
        no_sampling_count = 0
        no_measurement_count = 0
        candidate_count = 0

        for row_index, row in raw.iterrows():
            first_value = row.iloc[0]

            # Month headings in the source use labels such as "CY 2012 JANUARY".
            if isinstance(first_value, str):
                match = re.search(r"CY\s+(\d{4})\s+([A-Z]+)", first_value.upper())
                if match:
                    try:
                        current_month = pd.to_datetime(
                            f"{match.group(2).title()} 1 {match.group(1)}",
                            format="%B %d %Y"
                        )
                    except ValueError:
                        current_month = None
                    continue

            sample_number = pd.to_numeric(first_value, errors="coerce")
            location = row.iloc[1] if len(row) > 1 else None
            if pd.isna(sample_number) or pd.isna(location) or not str(location).strip():
                continue

            candidate_count += 1
            row_text = " ".join(str(value) for value in row.tolist()
                                if pd.notna(value)).upper()
            if "NO SAMPLING CONDUCTED" in row_text:
                no_sampling_total += 1
                no_sampling_count += 1
                continue

            # Source columns 6–17 contain the water-quality measurements.
            measurements = row.iloc[6:18].copy()
            measurements = measurements.replace(r"^\s*$", np.nan, regex=True)
            measurements = pd.to_numeric(measurements, errors="coerce")
            if measurements.notna().sum() == 0:
                no_measurement_count += 1
                continue

            # Use the explicit sample date when present; otherwise use the month heading.
            explicit_date = pd.to_datetime(row.iloc[4], errors="coerce")
            sample_date = explicit_date if pd.notna(explicit_date) else current_month
            clean_location = re.sub(r"\s+", " ", str(location)).strip().strip(",")

            record = {
                "Municipality": sheet,
                "Sample_No": int(sample_number),
                "Sampling_Location": clean_location,
                "Latitude_North_deg": pd.to_numeric(row.iloc[2], errors="coerce"),
                "Longitude_East_deg": pd.to_numeric(row.iloc[3], errors="coerce"),
                "Sample_Date": sample_date,
                "Sampling_Time": normalize_time(row.iloc[5]),
            }
            for position, column_name in enumerate(METRIC_NAMES[2:]):
                record[column_name] = measurements.iloc[position]
            records.append(record)
            usable_count += 1

        extraction_summary.append({
            "Sheet": sheet,
            "Candidate station rows": candidate_count,
            "Usable measurement rows": usable_count,
            "No sampling rows skipped": no_sampling_count,
            "Rows with no usable measurement skipped": no_measurement_count,
        })

    return pd.DataFrame(records), extraction_summary, no_sampling_total

def add_iqr_outlier_flag(data):
    """Flag IQR outliers for review; do not delete potentially genuine observations."""
    flag = pd.Series(False, index=data.index)
    counts = []
    for column in METRIC_NAMES[2:]:
        values = data[column].dropna()
        if values.empty:
            count = 0
        else:
            q1, q3 = values.quantile([0.25, 0.75])
            iqr = q3 - q1
            lower_fence = q1 - 1.5 * iqr
            upper_fence = q3 + 1.5 * iqr
            column_flag = (data[column] < lower_fence) | (data[column] > upper_fence)
            flag |= column_flag.fillna(False)
            count = int(column_flag.sum())
        counts.append({"Parameter": column, "IQR_Outlier_Count": count})
    data["Potential_Outlier_Flag"] = flag
    return data, pd.DataFrame(counts)

def format_workbook(file_path):
    workbook = load_workbook(file_path)
    for worksheet in workbook.worksheets:
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        for cell in worksheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="17365D")
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        worksheet.row_dimensions[1].height = 30
        for column in worksheet.columns:
            max_length = max(
                (len(str(cell.value)) for cell in list(column)[:150]
                 if cell.value is not None), default=10
            )
            worksheet.column_dimensions[
                get_column_letter(column[0].column)
            ].width = min(max(max_length + 2, 12), 42)
    workbook.save(file_path)

def main():
    data, extraction_summary, no_sampling_count = extract_records(INPUT_FILE)
    data["Sample_Date"] = pd.to_datetime(data["Sample_Date"], errors="coerce")
    # The duplicate key represents the same municipality/site/sample/date combination.
    duplicate_mask = data.duplicated(
        subset=["Municipality", "Sample_No", "Sampling_Location", "Sample_Date"],
        keep="first"
    )
    duplicates_removed = int(duplicate_mask.sum())
    data = data.loc[~duplicate_mask].copy()

    data, outlier_summary = add_iqr_outlier_flag(data)
    final_columns = [
        "Municipality", "Sample_No", "Sampling_Location", "Latitude_North_deg",
        "Longitude_East_deg", "Sample_Date", "Sampling_Time"
    ] + METRIC_NAMES[2:] + ["Potential_Outlier_Flag"]
    data = data[final_columns].sort_values(
        ["Municipality", "Sample_Date", "Sample_No"], na_position="last"
    ).reset_index(drop=True)

    quality_assessment = pd.DataFrame([
        {"Quality Check": "Header / report formatting",
         "Finding": "Source sheets contain report titles, guideline rows, month labels and unnamed columns.",
         "Action": "Extracted sample rows, assigned descriptive column names and parsed month/date values."},
        {"Quality Check": "No sampling conducted",
         "Finding": f"{no_sampling_count} rows explicitly marked NO SAMPLING CONDUCTED.",
         "Action": "Excluded because no water-quality measurements were recorded."},
        {"Quality Check": "Rows with no measurements",
         "Finding": f"{sum(x['Rows with no usable measurement skipped'] for x in extraction_summary)} station rows had no usable measurements.",
         "Action": "Excluded rows with no measurements; retained partially measured rows."},
        {"Quality Check": "Duplicate records",
         "Finding": f"{duplicates_removed} duplicates detected using municipality, sample number, location and date.",
         "Action": "Removed later duplicate occurrences based on that key."},
        {"Quality Check": "Missing values",
         "Finding": f"{int(data.isna().sum().sum())} missing cells remain in the final table.",
         "Action": "Left unknown measurements blank rather than inventing or imputing environmental readings."},
        {"Quality Check": "Data types",
         "Finding": "Coordinates and measurements are numeric; sample date is parsed as a date; time is standardized when available.",
         "Action": "Converted fields programmatically."},
        {"Quality Check": "Location formatting",
         "Finding": "Location labels contain inconsistent extra spaces.",
         "Action": "Trimmed and collapsed repeated whitespace; did not merge differently worded locations."},
        {"Quality Check": "Potential outliers",
         "Finding": f"{int(data['Potential_Outlier_Flag'].sum())} rows have at least one value outside the 1.5×IQR fences.",
         "Action": "Flagged for review but not removed because extreme values may be real observations."},
        {"Quality Check": "Unnecessary columns / report rows",
         "Finding": "Original workbook includes report-only rows and blank/formatted columns.",
         "Action": "Excluded report-only layout and unused columns from the analysis-ready dataset."},
    ])

    missing_summary = pd.DataFrame({
        "Column": data.columns,
        "Missing_Count": [int(data[column].isna().sum()) for column in data.columns],
        "Missing_Percent": [round(data[column].isna().mean() * 100, 2)
                            for column in data.columns],
    })
    overview = pd.DataFrame([
        ["Original workbook sheets", len(pd.ExcelFile(INPUT_FILE).sheet_names)],
        ["Usable measurement rows before duplicate removal", len(data) + duplicates_removed],
        ["Duplicate records removed", duplicates_removed],
        ["Final cleaned records", len(data)],
        ["Final variables", len(data.columns)],
        ["Rows flagged for potential outliers", int(data["Potential_Outlier_Flag"].sum())],
        ["Rows excluded: no sampling conducted", no_sampling_count],
        ["Missing cells in final dataset", int(data.isna().sum().sum())],
    ], columns=["Metric", "Value"])

    with pd.ExcelWriter(OUTPUT_FILE, engine="openpyxl") as writer:
        data.to_excel(writer, index=False, sheet_name="Cleaned_Data")
        overview.to_excel(writer, index=False, sheet_name="Summary")
        quality_assessment.to_excel(writer, index=False, sheet_name="Quality_Assessment")
        missing_summary.to_excel(writer, index=False, sheet_name="Missing_Values")
        pd.DataFrame(extraction_summary).to_excel(
            writer, index=False, sheet_name="Sheet_Extraction"
        )
        outlier_summary.to_excel(writer, index=False, sheet_name="Outlier_Summary")

    format_workbook(OUTPUT_FILE)
    print(f"Preprocessing complete: {OUTPUT_FILE}")
    print(f"Final records: {len(data)} | Variables: {len(data.columns)}")
    print(f"Rows flagged for outlier review: {int(data['Potential_Outlier_Flag'].sum())}")

if __name__ == "__main__":
    main()
